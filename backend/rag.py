import asyncio
import hashlib
import logging
import time
import json
import io
import re
import uuid
import chromadb
import httpx
import ollama
from pathlib import Path, PurePosixPath
from typing import Optional
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.embeddings import Embeddings
from langchain_ollama import OllamaEmbeddings
from langchain_community.vectorstores import Chroma

import config

log = logging.getLogger(__name__)

chroma_client = chromadb.PersistentClient(path=config.CHROMA_PERSIST_DIR)
collection = chroma_client.get_or_create_collection(name=config.CHROMA_COLLECTION)



class GeminiEmbeddings(Embeddings):
    """Hosted embeddings through Gemini's OpenAI-compatible endpoint, so no Ollama server is needed."""

    def __init__(self, model: str, api_key: str, base_url: str, batch_size: int = 50, max_retries: int = 3):
        self.model = model
        self.api_key = api_key
        self.url = base_url.rstrip("/") + "/embeddings"
        self.batch_size = max(1, batch_size)
        self.max_retries = max_retries

    def _embed(self, texts: list[str]) -> list[list[float]]:
        if not self.api_key:
            raise RuntimeError("GEMINI_API_KEY is not set on the server, so uploads can't be indexed")
        vectors: list[list[float]] = []
        with httpx.Client(timeout=60) as client:
            for start in range(0, len(texts), self.batch_size):
                batch = texts[start:start + self.batch_size]
                for attempt in range(self.max_retries + 1):
                    response = client.post(
                        self.url,
                        headers={"Authorization": f"Bearer {self.api_key}"},
                        json={"model": self.model, "input": batch},
                    )
                    # rate limit / temporary server trouble: wait and try again
                    if response.status_code in (429, 500, 502, 503) and attempt < self.max_retries:
                        time.sleep(2 ** attempt)
                        continue
                    break
                if response.status_code >= 400:
                    raise RuntimeError(f"Embedding request failed ({response.status_code}): {response.text[:200]}")
                rows = sorted(response.json()["data"], key=lambda row: row.get("index", 0))
                vectors.extend(row["embedding"] for row in rows)
        return vectors

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return self._embed(texts)

    def embed_query(self, text: str) -> list[float]:
        return self._embed([text])[0]


if config.EMBEDDING_PROVIDER == "gemini":
    embeddings = GeminiEmbeddings(
        config.EMBEDDING_MODEL, config.GEMINI_API_KEY, config.GEMINI_BASE_URL, config.EMBEDDING_BATCH_SIZE,
    )
else:
    embeddings = OllamaEmbeddings(model="nomic-embed-text", base_url=config.OLLAMA_BASE_URL)
_async_ollama_client = ollama.AsyncClient(host=config.OLLAMA_BASE_URL)
_chat_semaphore = asyncio.Semaphore(config.MAX_CONCURRENT_CHATS)
_waiting_count = 0
_waiting_lock = asyncio.Lock()


# ---------------------------------------------------------------- Utils
def _image_data_url(data: str) -> str:
    if data.startswith("data:"):
        return data
    mime = "image/png" if data.startswith("iVBOR") else "image/webp" if data.startswith("UklGR") else "image/jpeg"
    return f"data:{mime};base64,{data}"

def _to_openai_messages(messages: list[dict], image_index: int | None = None) -> tuple[list[dict], bool]:
    out = []
    has_images = image_index is not None and bool(messages[image_index].get("images"))
    for index, m in enumerate(messages):
        if has_images and index == image_index:
            parts = [{"type": "text", "text": m["content"]}]
            parts += [{"type": "image_url", "image_url": {"url": _image_data_url(i)}} for i in m["images"]]
            out.append({"role": m["role"], "content": parts})
        else:
            out.append({"role": m["role"], "content": m["content"]})
    return out, has_images


# ---------------------------------------------------------------- Groq
def _use_groq(messages: list[dict], image_index: int | None = None) -> bool:
    return bool(config.GROQ_API_KEY) and (bool(config.GROQ_VISION_MODEL) or image_index is None)

def _groq_request(messages: list[dict], stream: bool, image_index: int | None = None) -> tuple[str, dict, dict]:
    msgs, has_images = _to_openai_messages(messages, image_index)
    body = {
        "model": config.GROQ_VISION_MODEL if has_images else config.GROQ_MODEL,
        "messages": msgs,
        "max_tokens": config.GROQ_MAX_TOKENS,
        "stream": stream,
    }
    if body["model"].startswith("openai/gpt-oss") and config.GROQ_REASONING_EFFORT:
        body["reasoning_effort"] = config.GROQ_REASONING_EFFORT
    headers = {"Authorization": f"Bearer {config.GROQ_API_KEY}"}
    return f"{config.GROQ_BASE_URL.rstrip('/')}/chat/completions", headers, body

def _raise_groq_error(response: httpx.Response) -> None:
    if response.status_code < 400:
        return
    try:
        detail = response.json()["error"]["message"]
    except Exception:
        detail = response.text[:300]
    raise RuntimeError(f"Groq error {response.status_code}: {detail}")

async def _groq_retry_wait(response: httpx.Response, attempt: int) -> bool:
    if response.status_code not in (429, 503) or attempt >= config.GROQ_MAX_RETRIES:
        return False
    try:
        wait = float(response.headers.get("retry-after", ""))
    except ValueError:
        wait = 2 ** attempt
    print(f"[groq] {response.status_code}; retrying in {wait:.0f}s", flush=True)
    await asyncio.sleep(min(wait, 30))
    return True

async def _groq_stream(messages: list[dict], image_index: int | None = None):
    url, headers, body = _groq_request(messages, stream=True, image_index=image_index)
    async with httpx.AsyncClient(timeout=120) as client:
        for attempt in range(config.GROQ_MAX_RETRIES + 1):
            async with client.stream("POST", url, headers=headers, json=body) as response:
                if response.status_code >= 400:
                    await response.aread()
                    if await _groq_retry_wait(response, attempt):
                        continue
                    _raise_groq_error(response)
                async for line in response.aiter_lines():
                    if not line.startswith("data:"):
                        continue
                    payload = line[5:].strip()
                    if payload == "[DONE]":
                        return
                    choice = json.loads(payload)["choices"][0]
                    yield choice["delta"].get("content") or "", choice.get("finish_reason")
                return

async def _groq_generate(messages: list[dict]) -> str:
    url, headers, body = _groq_request(messages, stream=False)
    async with httpx.AsyncClient(timeout=120) as client:
        for attempt in range(config.GROQ_MAX_RETRIES + 1):
            response = await client.post(url, headers=headers, json=body)
            if await _groq_retry_wait(response, attempt):
                continue
            _raise_groq_error(response)
            return response.json()["choices"][0]["message"]["content"].strip()


# ---------------------------------------------------------------- Gemini
def _gemini_request(messages: list[dict], stream: bool, image_index: int | None = None) -> tuple[str, dict, dict]:
    msgs, _ = _to_openai_messages(messages, image_index)
    body = {
        "model": config.GEMINI_MODEL,
        "messages": msgs,
        "max_tokens": config.GEMINI_MAX_TOKENS,
        "stream": stream,
    }
    headers = {"Authorization": f"Bearer {config.GEMINI_API_KEY}"}
    return f"{config.GEMINI_BASE_URL.rstrip('/')}/chat/completions", headers, body

async def _gemini_retry_wait(response: httpx.Response, attempt: int) -> bool:
    if response.status_code not in (429, 503) or attempt >= config.GEMINI_MAX_RETRIES:
        return False
    try:
        wait = float(response.headers.get("retry-after", ""))
    except ValueError:
        wait = 2 ** attempt
    print(f"[gemini] {response.status_code}; retrying in {wait:.0f}s", flush=True)
    await asyncio.sleep(min(wait, 30))
    return True

async def _gemini_stream(messages: list[dict], image_index: int | None = None):
    url, headers, body = _gemini_request(messages, stream=True, image_index=image_index)
    async with httpx.AsyncClient(timeout=120) as client:
        for attempt in range(config.GEMINI_MAX_RETRIES + 1):
            async with client.stream("POST", url, headers=headers, json=body) as response:
                if response.status_code >= 400:
                    await response.aread()
                    if await _gemini_retry_wait(response, attempt):
                        continue
                    raise RuntimeError(f"Gemini error {response.status_code}: {response.text[:300]}")
                
                async for line in response.aiter_lines():
                    if not line.startswith("data:"):
                        continue
                    payload = line[5:].strip()
                    if payload == "[DONE]":
                        return
                    try:
                        choice = json.loads(payload)["choices"][0]
                        yield choice["delta"].get("content") or "", choice.get("finish_reason")
                    except (KeyError, json.JSONDecodeError):
                        continue
                return

async def _gemini_generate(messages: list[dict]) -> str:
    url, headers, body = _gemini_request(messages, stream=False)
    async with httpx.AsyncClient(timeout=120) as client:
        for attempt in range(config.GEMINI_MAX_RETRIES + 1):
            response = await client.post(url, headers=headers, json=body)
            if await _gemini_retry_wait(response, attempt):
                continue
            if response.status_code >= 400:
                raise RuntimeError(f"Gemini error {response.status_code}: {response.text[:300]}")
            return response.json()["choices"][0]["message"]["content"].strip()


# ---------------------------------------------------------------- Core RAG & Fallback Logic
async def _answer_source(messages: list[dict], options: dict, image_index: int | None):
    """Yields (text, finish_reason). Attempts Groq -> Gemini -> Ollama fallback chain."""
    started = False

    if _use_groq(messages, image_index):
        try:
            async for delta, reason in _groq_stream(messages, image_index):
                started = started or bool(delta)
                yield delta, reason
            return
        except RuntimeError as error:
            if started:
                raise # Cannot fallback cleanly if tokens are already streaming to the frontend
            print(f"[groq] request failed ({error}); falling back to next provider...", flush=True)

    if config.GEMINI_API_KEY:
        try:
            async for delta, reason in _gemini_stream(messages, image_index):
                started = started or bool(delta)
                yield delta, reason
            return
        except RuntimeError as error:
            if started:
                raise
            print(f"[gemini] request failed ({error}); falling back to local Ollama...", flush=True)

    # Final fallback if both external providers fail or are not configured
    async for item in _ollama_stream(messages, options):
        yield item

async def _ollama_stream(messages: list[dict], options: dict):
    stream = await _async_ollama_client.chat(
        model=config.OLLAMA_MODEL, messages=messages, think=False, stream=True, options=options,
    )
    async for chunk in stream:
        yield chunk["message"]["content"], chunk.get("done_reason")

def _table_to_markdown(table: list[list]) -> str:
    rows = [["" if cell is None else str(cell).strip() for cell in row] for row in table if row]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)

def extract_text_by_page(file_path: str) -> list[tuple[int, str, list[str]]]:
    path = Path(file_path)
    return extract_text_from_bytes(path.name, path.read_bytes())

def extract_text_from_bytes(filename: str, contents: bytes) -> list[tuple[int, str, list[str]]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import pdfplumber
        pages = []
        with pdfplumber.open(io.BytesIO(contents)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                tables = [_table_to_markdown(t) for t in page.extract_tables()]
                tables = [t for t in tables if t]
                if text or tables:
                    pages.append((i, text, tables))
        return pages

    elif suffix in {".txt", ".md"}:
        return [(0, contents.decode("utf-8-sig"), [])]

    elif suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(contents))
        text = ""
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        tables = []
        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            md = _table_to_markdown(rows)
            if md:
                tables.append(md)
        return [(0, text, tables)]

    raise ValueError(f"Unsupported material file type: {suffix or '(no extension)'}")

def _index_files(vault_files: list[Path], vault_dir: Path, branch: str = "General") -> tuple[int, set[str]]:
    """Index files that sit in the local uploads folder."""
    documents = []
    for vault_file in vault_files:
        rel_path = vault_file.relative_to(vault_dir)
        note_branch = rel_path.parts[1] if len(rel_path.parts) > 2 else branch
        documents.append((rel_path.as_posix(), note_branch, vault_file.name, vault_file.read_bytes()))
    return _index_documents(documents)


def _index_documents(documents: list[tuple[str, str, str, bytes]]) -> tuple[int, set[str]]:
    """Chunk, embed and store files given as (source, branch, file name, bytes). `source` is the stable key
    the chunks are filed under (and later searched by); it is the same string whether the original lives
    on disk or in Google Drive."""
    vault_collection = chroma_client.get_or_create_collection(name=config.CHROMA_VAULT_COLLECTION)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000, chunk_overlap=100,
        separators=["\n\n", "\n", " ", ""]
    )

    all_chunks = []
    all_metadatas = []
    all_ids = []

    for source, note_branch, file_name, contents in documents:
        path_hash = hashlib.sha256(source.encode("utf-8")).hexdigest()
        note_type = "uploaded_material"
        pages = extract_text_from_bytes(file_name, contents)

        for page_number, page_text, tables in pages:
            content = "\n\n".join([page_text, *tables]).strip()
            if not content:
                continue
            for j, chunk in enumerate(splitter.split_text(content)):
                all_chunks.append(chunk)
                all_metadatas.append({
                    "source": source,
                    "page": page_number,
                    "branch": note_branch,
                    "type": note_type,
                })
                all_ids.append(f"obsidian_{path_hash}_page{page_number}_part{j}")

    if all_chunks:
        vault_collection.upsert(
            documents=all_chunks,
            metadatas=all_metadatas,
            ids=all_ids,
            embeddings=embeddings.embed_documents(all_chunks)
        )

    return len(all_chunks), set(all_ids)

def _safe_branch(branch: str) -> str:
    return re.sub(r"[^A-Za-z0-9 _-]", "", branch).strip() or "General"


def _prepare_uploads(files: list[tuple[str, bytes]], storage_ids: Optional[list[str]] = None) -> list[tuple[str, bytes, str]]:
    """Validate uploads (name, type, readable text) before anything is stored anywhere.
    Returns (safe file name, bytes, storage id) for each; raises ValueError for a bad file."""
    if storage_ids is not None and len(storage_ids) != len(files):
        raise ValueError("Each uploaded material must have one storage ID")
    prepared_files = []

    for index, (filename, contents) in enumerate(files):
        display_name = Path(filename.replace("\\", "/")).name
        safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", display_name).strip(" .")
        if not safe_name:
            raise ValueError("Uploaded material must have a valid filename")
        suffix = Path(safe_name).suffix.lower()
        if suffix not in {".pdf", ".docx", ".txt", ".md"}:
            raise ValueError(f"Unsupported material file type: {suffix or '(no extension)'}")
        safe_name = f"{Path(safe_name).stem[:100]}{suffix}"
        pages = extract_text_from_bytes(safe_name, contents)
        if not any(text.strip() or any(table.strip() for table in tables) for _, text, tables in pages):
            raise ValueError(f"No readable text found in uploaded file: {display_name}")
        storage_id = storage_ids[index] if storage_ids is not None else uuid.uuid4().hex
        prepared_files.append((safe_name, contents, storage_id))
    return prepared_files


def store_uploaded_materials(
    files: list[tuple[str, bytes]],
    branch: str,
    storage_ids: Optional[list[str]] = None,
) -> list[str]:
    """Local-disk storage: write validated uploads into the uploads folder."""
    vault_dir = Path(config.CHROMA_VAULT_DIR).expanduser()
    vault_dir.mkdir(parents=True, exist_ok=True)
    prepared_files = _prepare_uploads(files, storage_ids)
    output_dir = vault_dir / "Uploaded Materials" / _safe_branch(branch)
    stored_paths = []

    output_dir.mkdir(parents=True, exist_ok=True)
    for safe_name, contents, storage_id in prepared_files:
        material_path = output_dir / f"{storage_id}_{safe_name}"
        material_path.write_bytes(contents)
        stored_paths.append(material_path.relative_to(vault_dir).as_posix())

    return stored_paths

class StoredPath(str):
    """The name an upload is filed under (what `Material.filename` holds). It is a plain str, so every
    caller that only wants the name keeps working; `drive_file_id` is set when the original went to Drive."""
    drive_file_id: Optional[str] = None

    def __new__(cls, path: str, drive_file_id: Optional[str] = None):
        obj = super().__new__(cls, path)
        obj.drive_file_id = drive_file_id
        return obj


def _ingest_to_drive(files: list[tuple[str, bytes]], branch: str, storage_ids: Optional[list[str]]) -> list[StoredPath]:
    """Validate, put the originals in Google Drive, then index them. Either everything lands or nothing does:
    if indexing (or a later upload) fails, the Drive files already uploaded are trashed again."""
    import drive_storage

    prepared = _prepare_uploads(files, storage_ids)
    safe_branch = _safe_branch(branch)
    uploaded: list[tuple[StoredPath, str, bytes]] = []
    try:
        for safe_name, contents, storage_id in prepared:
            drive_name = f"{storage_id}_{safe_name}"
            file_id = drive_storage.upload_bytes(drive_name, contents)
            uploaded.append((StoredPath(f"Uploaded Materials/{safe_branch}/{drive_name}", file_id), safe_name, contents))
        _index_documents([(str(stored), safe_branch, safe_name, contents) for stored, safe_name, contents in uploaded])
    except Exception:
        for stored, _, _ in uploaded:
            try:
                drive_storage.delete_file(stored.drive_file_id)
            except Exception:
                log.warning("Couldn't remove Drive file %s after a failed upload", stored.drive_file_id)
        raise
    return [stored for stored, _, _ in uploaded]


def sync_drive_index(materials: list[tuple[str, str, str]], force: bool = False) -> dict:
    """Make the vector store match the files in Drive. `materials` is (filename, branch, drive file id) for
    every Material row stored in Drive. Files with no chunks yet are downloaded and indexed (all of them
    after a Render restart, which wipes the index); with `force`, every file is re-indexed. Chunks of files
    that no longer have a row are dropped. A file that can't be fetched is skipped and reported, not fatal."""
    import drive_storage

    vault_collection = chroma_client.get_or_create_collection(name=config.CHROMA_VAULT_COLLECTION)
    wanted = {_material_index_source(filename) for filename, _, _ in materials}
    chunks = 0
    indexed_files = 0
    failed: list[str] = []

    for filename, branch, file_id in materials:
        source = _material_index_source(filename)
        if not force and vault_collection.get(where={"source": source}, limit=1)["ids"]:
            continue
        try:
            contents = drive_storage.download_bytes(file_id)
            count, _ = _index_documents([(source, branch, PurePosixPath(source).name, contents)])
            chunks += count
            indexed_files += 1
        except Exception as exc:
            log.warning("Couldn't index %s from Drive: %s", filename, exc)
            failed.append(filename)

    existing = vault_collection.get(where={"type": "uploaded_material"}, include=["metadatas"])
    stale = [
        doc_id for doc_id, metadata in zip(existing["ids"], existing["metadatas"])
        if (metadata or {}).get("source") not in wanted
    ]
    if stale:
        vault_collection.delete(ids=stale)
    return {"chunks": chunks, "indexed_files": indexed_files, "failed": failed, "removed_chunks": len(stale)}


def ingest_uploaded_materials(
    files: list[tuple[str, bytes]],
    branch: str = "General",
    storage_ids: Optional[list[str]] = None,
) -> list[str]:
    import drive_storage
    if drive_storage.enabled():
        return _ingest_to_drive(files, branch, storage_ids)
    stored_paths = store_uploaded_materials(files, branch, storage_ids)
    vault_dir = Path(config.CHROMA_VAULT_DIR).expanduser()
    material_paths = [vault_dir / PurePosixPath(path) for path in stored_paths]
    try:
        _index_files(
            material_paths,
            vault_dir,
            branch,
        )
    except Exception:
        for path in material_paths:   # no half-saved uploads: a file with no index and no DB record is an orphan
            path.unlink(missing_ok=True)
        raise
    return stored_paths

def ingest_uploads(uploads_path: str, branch: str = "General"):
    vault_collection = chroma_client.get_or_create_collection(name=config.CHROMA_VAULT_COLLECTION)
    vault_dir = Path(uploads_path).expanduser()
    vault_dir.mkdir(parents=True, exist_ok=True)
    materials_dir = vault_dir / "Uploaded Materials"
    vault_files = [
        path for path in (materials_dir.rglob("*") if materials_dir.is_dir() else [])
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}
    ]
    chunk_count, current_ids = _index_files(sorted(vault_files), vault_dir, branch)

    existing = vault_collection.get(
        include=["metadatas"],
    )
    managed_types = {"obsidian_note", "uploaded_material"}
    stale_ids = sorted(
        doc_id
        for doc_id, metadata in zip(existing["ids"], existing["metadatas"])
        if (metadata or {}).get("type") in managed_types and doc_id not in current_ids
    )
    if stale_ids:
        vault_collection.delete(ids=stale_ids)

    return chunk_count


_PAGE_RE = re.compile(r"pages?\s*(\d+)(?:\s*(?:-|–|to|and)\s*(\d+))?", re.IGNORECASE)

def _page_request(question: str):
    m = _PAGE_RE.search(question)
    if not m:
        return None
    start = int(m.group(1))
    end = int(m.group(2)) if m.group(2) else start
    if end < start:
        start, end = end, start
    return (start, end)

def _material_index_source(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).as_posix()


def delete_material_chunks(filename: str) -> int:
    """Remove the AI's chunks for one uploaded file (every page). Returns how many were removed.
    (admin_routes.delete_file already called this, but it was never defined, so deleting a file always failed.)"""
    vault_collection = chroma_client.get_or_create_collection(name=config.CHROMA_VAULT_COLLECTION)
    ids = vault_collection.get(where={"source": _material_index_source(filename)})["ids"]
    if ids:
        vault_collection.delete(ids=ids)
    return len(ids)

def _chunks_for_pages(source: str, start: int, end: int, limit: int = 8):
    vector_store = Chroma(
        collection_name=config.CHROMA_VAULT_COLLECTION,
        embedding_function=embeddings,
        client=chroma_client,
    )
    got = vector_store._collection.get(where={"source": source}, include=["documents", "metadatas"])
    matches = []
    for text, meta in zip(got["documents"], got["metadatas"]):
        page = (meta or {}).get("page")
        if page is not None and start <= page <= end:
            matches.append((page, text))
    matches.sort(key=lambda t: t[0])
    return matches[:limit]   

async def ask_stream(
    question: str,
    chat_history: list[dict] = None,
    material_filename: str = None,
    images: list[str] = None,
):
    context = ""
    if material_filename:
        material_source = _material_index_source(material_filename)
        page_range = _page_request(question)
        if page_range:
            page_chunks = await asyncio.to_thread(_chunks_for_pages, material_source, *page_range)
            context = "\n\n".join(f"[Page {p}] {t}" for p, t in page_chunks)
        else:
            results = await asyncio.to_thread(_similarity_search, question, material_source)
            context = "\n\n".join(
                f"[Page {d.metadata.get('page', '?')}] {d.page_content}" for d in results
            )
    else:
        results = await asyncio.to_thread(_similarity_search, question)
        context = "\n\n".join(d.page_content for d in results)

    system_prompt = (
        "You are a friendly math assistant for students. "
        "You can have normal conversation — greetings, small talk, thanking you, "
        "asking who/what you are — respond naturally and briefly to those. "
        "For anything that is actually a question or request for help, ONLY help if "
        "it is related to mathematics (arithmetic, algebra, geometry, statistics, "
        "calculus, etc.). If someone asks for substantive help on a non-math topic "
        "(history, coding, writing, other homework, etc.), politely decline and say: "
        "'I can only help with math questions. Please ask me something related to mathematics.' "
        "For math problems, show your work step by step. "
        "Be clear, concise, and structured in your explanations. "
        "Never use emojis or emoticons; use plain text only. "
        "If reference material is provided below, it comes from course documents "
        "in a shared knowledge base for THIS question only — "
        "it is not something you said earlier, and you have no memory of using or not "
        "using it on any previous turn. Use it only if it genuinely helps answer the "
        "current question; if it doesn't, ignore it silently and never apologize for "
        "'not referencing' something from a prior answer. "
        "If you don't know the answer, say so honestly. "
        "You are able to see and understand images the student attaches — photos of "
        "handwritten work, diagrams, charts, textbook pages. When an image is attached, "
        "actually look at it and describe or work with what's specifically in it; "
        "never say you can't see images. "
        "When the student sends an image together with a message, treat the message as their "
        "instruction about the image (for example 'solve number 5', 'check my work', 'explain this graph'). "
        "If they name a specific item, work only on that. If the message is vague, work out what the image "
        "shows (a problem to solve, work to check, a diagram to explain), do the most likely helpful thing, "
        "and say in one short line what you assumed. When checking handwritten work, first state what you "
        "read for each step; if a symbol or number is unclear, say so instead of guessing. "
        "If you cannot read the image or cannot tell which problem they mean, say what you can see and ask one short question."
    )
    if material_filename:
        display = re.sub(
            r"^[a-f0-9]{32}_",
            "",
            PurePosixPath(material_filename.replace("\\", "/")).name,
            flags=re.IGNORECASE,
        )
        system_prompt += (
            f"\n\nThe student currently has the document \"{display}\" open and is asking about it. "
            "The excerpts below are taken only from that document, labelled with their page numbers; "
            "when you use them, mention the page. If the excerpts don't contain what the student "
            "asks about, say you couldn't find it in the parts of the document you can see, "
            "rather than guessing."
        )
        if context:
            system_prompt += f"\n\nExcerpts from the open document:\n{context}"
        else:
            system_prompt += "\n\n(No excerpts from the open document matched this question.)"
    elif context:
        system_prompt += f"\n\nReference material for this question only (use if relevant, otherwise ignore):\n{context}"

    messages = [{"role": "system", "content": system_prompt}]

    if chat_history:
        recent_history = chat_history[-10:]
        image_message_indexes = [
            index for index, msg in enumerate(recent_history) if msg.get("images")
        ]
        keep_image_indexes = (
            set(image_message_indexes[-config.MAX_HISTORY_IMAGES:])
            if config.MAX_HISTORY_IMAGES > 0
            else set()
        )
        for index, msg in enumerate(recent_history):
            history_message = {"role": msg["role"], "content": msg["content"]}
            if index in keep_image_indexes:
                history_images = msg["images"]
                history_message["images"] = [
                    image["data"] if isinstance(image, dict) else image
                    for image in history_images
                ]
            messages.append(history_message)

    user_message = {"role": "user", "content": question}
    image_index = None
    if images:
        user_message["images"] = images
        image_index = len(messages)
    messages.append(user_message)

    global _waiting_count
    async with _waiting_lock:
        _waiting_count += 1
        queue_position = _waiting_count

    yield {"type": "meta", "was_queued": queue_position > 1, "queue_position": queue_position}

    async with _chat_semaphore:
        async with _waiting_lock:
            _waiting_count -= 1
        full_answer = []
        options = {"num_ctx": config.OLLAMA_NUM_CTX, "num_predict": -1}
        for round_number in range(config.OLLAMA_MAX_CONTINUATIONS + 1):
            round_text = []
            done_reason = None
            source = _answer_source(messages, options, image_index)
            async for delta, reason in source:
                if delta:
                    round_text.append(delta)
                    full_answer.append(delta)
                    yield {"type": "token", "text": delta}
                if reason:
                    done_reason = reason

            if done_reason != "length" or not round_text:
                break

            print(
                f"[rag] answer hit a length limit; continuing (round {round_number + 1})",
                flush=True,
            )
            messages = messages + [
                {"role": "assistant", "content": "".join(round_text)},
                {
                    "role": "user",
                    "content": "Continue exactly where you left off. Do not repeat anything or start over.",
                },
            ]

    yield {"type": "done", "answer": "".join(full_answer)}   

def queue_status() -> dict:
    return {"waiting": _waiting_count, "max_concurrent": config.MAX_CONCURRENT_CHATS}

def _similarity_search(question: str, source_filter: str = None):
    vector_store = Chroma(
        collection_name=config.CHROMA_VAULT_COLLECTION,
        embedding_function=embeddings,
        client=chroma_client,
    )
    kwargs = {"filter": {"source": source_filter}} if source_filter else {}
    k = 5 if source_filter else 3
    results = vector_store.similarity_search_with_score(question, k=k, **kwargs)
    docs = [doc for doc, distance in results if distance <= config.MAX_CONTEXT_DISTANCE]
    return docs

def _material_context(query: str, branch: str, k: int = 6) -> str:
    store = Chroma(
        collection_name=config.CHROMA_VAULT_COLLECTION,
        embedding_function=embeddings,
        client=chroma_client,
    )
    hits = store.similarity_search_with_score(query, k=k, filter={"branch": branch})
    return "\n\n".join(d.page_content for d, distance in hits if distance <= config.MAX_CONTEXT_DISTANCE)

async def generate_text(system: str, prompt: str, branch: str, query: str) -> str:
    context = await asyncio.to_thread(_material_context, query, branch) if branch else ""
    if context:
        system += f"\n\nCourse material to base your answer on (ignore parts that don't fit):\n{context}"
    messages = [{"role": "system", "content": system}, {"role": "user", "content": prompt}]
    
    async with _chat_semaphore:
        # 1. Try Groq
        if config.GROQ_API_KEY:
            try:
                return await _groq_generate(messages)
            except RuntimeError as error:
                print(f"[groq] generation failed ({error}); falling back...", flush=True)

        # 2. Try Gemini
        if config.GEMINI_API_KEY:
            try:
                return await _gemini_generate(messages)
            except RuntimeError as error:
                print(f"[gemini] generation failed ({error}); falling back...", flush=True)

        # 3. Try Hermes
        if config.HERMES_API_URL:
            async with httpx.AsyncClient(timeout=config.HERMES_TIMEOUT) as client:
                response = await client.post(
                    f"{config.HERMES_API_URL.rstrip('/')}/chat/completions",
                    headers={"Authorization": f"Bearer {config.HERMES_API_KEY}"},
                    json={"model": "hermes-agent", "messages": messages},
                )
                response.raise_for_status()
                return response.json()["choices"][0]["message"]["content"].strip()
                
        # 4. Fallback to Local Ollama
        response = await _async_ollama_client.chat(
            model=config.OLLAMA_MODEL,
            messages=messages,
            think=False,
            options={"num_ctx": config.OLLAMA_NUM_CTX, "num_predict": -1},
        )
    return response["message"]["content"].strip()